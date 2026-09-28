#!/usr/bin/env python3
"""Concurrent, same-site SEO crawler that writes a formatted Excel audit report."""
from __future__ import annotations

import argparse
import hashlib
import logging
import re
import sys
import threading
import time
from collections import Counter, defaultdict, deque
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable
from urllib.parse import urldefrag, urljoin, urlparse, urlunparse
from urllib.robotparser import RobotFileParser

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from tqdm import tqdm

class SafeStreamHandler(logging.StreamHandler):
    def emit(self, record):
        try:
            super().emit(record)
        except Exception:
            pass

logging.basicConfig(level=logging.INFO, handlers=[SafeStreamHandler()])
logging.getLogger("urllib3").setLevel(logging.CRITICAL)
logging.getLogger("requests").setLevel(logging.CRITICAL)

USER_AGENT = "SEO-Site-Auditor/1.0 (+https://example.invalid)"
REQUEST_TIMEOUT = 20
HTML_TYPES = ("text/html", "application/xhtml+xml")
session = requests.Session()
session.headers.update({"User-Agent": USER_AGENT, "Accept": "text/html,application/xhtml+xml,*/*;q=0.8"})
thread_local = threading.local()


@dataclass
class Resource:
    url: str
    alt: str = ""
    anchor: str = ""


@dataclass
class Page:
    url: str
    status: int | None = None
    response_ms: int | None = None
    redirect_chain: str = ""
    title: str = ""
    meta_description: str = ""
    h1s: list[str] = field(default_factory=list)
    word_count: int = 0
    canonical: str = ""
    robots: str = ""
    content_hash: str = ""
    internal_links: list[Resource] = field(default_factory=list)
    images: list[Resource] = field(default_factory=list)
    assets: list[Resource] = field(default_factory=list)
    error: str = ""

    @property
    def indexability(self) -> str:
        return "Noindex" if "noindex" in self.robots.lower() else "Index"


def normalise_url(url: str) -> str:
    """Remove fragments and normalise a URL without changing meaningful queries."""
    url = urldefrag(url)[0].strip()
    parts = urlparse(url)
    path = parts.path or "/"
    return urlunparse((parts.scheme.lower(), parts.netloc.lower(), path, "", parts.query, ""))


def get_session() -> requests.Session:
    if not hasattr(thread_local, "session"):
        thread_local.session = requests.Session()
        thread_local.session.headers.update(session.headers)
    return thread_local.session


def fetch(url: str, stream: bool = False) -> requests.Response:
    return get_session().get(url, timeout=REQUEST_TIMEOUT, allow_redirects=True, stream=stream)


def same_domain(url: str, root: str, include_subdomains: bool) -> bool:
    host, root_host = urlparse(url).hostname or "", urlparse(root).hostname or ""
    return host == root_host or (include_subdomains and host.endswith("." + root_host))


def load_robots(root_url: str, ignore: bool) -> RobotFileParser | None:
    if ignore:
        return None
    robots_url = urljoin(root_url, "/robots.txt")
    parser = RobotFileParser()
    parser.set_url(robots_url)
    try:
        parser.read()
        return parser
    except Exception as exc:
        logging.warning("Could not read robots.txt (%s): %s", robots_url, exc)
        return None  # Crawl continues if robots cannot be retrieved.


def allowed(url: str, robots: RobotFileParser | None) -> bool:
    return robots is None or robots.can_fetch(USER_AGENT, url)


def text_content(soup: BeautifulSoup) -> str:
    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()
    return re.sub(r"\s+", " ", soup.get_text(" ", strip=True)).strip()


def analyse_page(url: str) -> Page:
    """Fetch one page and extract its on-page data. Network errors become Page errors."""
    page = Page(url=url)
    try:
        started = time.perf_counter()
        response = fetch(url)
        page.response_ms = round((time.perf_counter() - started) * 1000)
        page.status = response.status_code
        page.redirect_chain = " -> ".join([r.url for r in response.history] + [response.url]) if response.history else ""
        if response.status_code >= 400 or not any(t in response.headers.get("Content-Type", "").lower() for t in HTML_TYPES):
            return page
        soup = BeautifulSoup(response.content, "lxml")
        title = soup.find("title")
        description = soup.find("meta", attrs={"name": re.compile(r"^description$", re.I)})
        robots = soup.find("meta", attrs={"name": re.compile(r"^robots$", re.I)})
        canonical = soup.find("link", attrs={"rel": lambda r: r and "canonical" in [x.lower() for x in (r if isinstance(r, list) else [r])]})
        page.title = title.get_text(" ", strip=True) if title else ""
        page.meta_description = description.get("content", "").strip() if description else ""
        page.robots = robots.get("content", "").strip() if robots else ""
        page.canonical = normalise_url(urljoin(response.url, canonical.get("href"))) if canonical and canonical.get("href") else ""
        page.h1s = [h.get_text(" ", strip=True) for h in soup.find_all("h1")]
        body_text = text_content(soup)
        page.word_count = len(re.findall(r"\b\w+[\w'-]*\b", body_text))
        page.content_hash = hashlib.sha1(body_text.lower().encode("utf-8")).hexdigest() if body_text else ""
        for a in soup.find_all("a", href=True):
            href = normalise_url(urljoin(response.url, a["href"]))
            if urlparse(href).scheme in ("http", "https"):
                page.internal_links.append(Resource(href, anchor=a.get_text(" ", strip=True)))
        for image in soup.find_all("img", src=True):
            page.images.append(Resource(normalise_url(urljoin(response.url, image["src"])), alt=image.get("alt", "").strip()))
        for link in soup.find_all("link", href=True):
            rel = " ".join(link.get("rel", [])).lower()
            if "stylesheet" in rel:
                page.assets.append(Resource(normalise_url(urljoin(response.url, link["href"])), alt="CSS"))
        for script in soup.find_all("script", src=True):
            page.assets.append(Resource(normalise_url(urljoin(response.url, script["src"])), alt="JS"))
    except requests.RequestException as exc:
        page.error = str(exc)
        try:
            logging.error("%s: %s", url, exc)
        except Exception:
            pass
    except Exception as exc:  # Parsing faults must never stop crawl.
        page.error = str(exc)
        try:
            logging.exception("Unexpected error on %s", url)
        except Exception:
            pass
    return page


def safe_is_tty() -> bool:
    try:
        return bool(sys.stderr and hasattr(sys.stderr, "isatty") and sys.stderr.isatty())
    except Exception:
        return False


def extract_sitemap_urls(root_url: str, max_urls: int = 10000) -> list[str]:
    """Helper to discover all URLs listed in sitemaps to seed crawler queue."""
    urls = []
    pending = deque([urljoin(root_url, "/sitemap.xml"), urljoin(root_url, "/sitemap_index.xml")])
    visited = set()
    while pending and len(urls) < max_urls:
        sitemap = pending.popleft()
        if sitemap in visited:
            continue
        visited.add(sitemap)
        try:
            r = fetch(sitemap)
            if r.status_code >= 400:
                continue
            soup = BeautifulSoup(r.content, "xml")
            for node in soup.find_all("sitemap"):
                loc = node.find("loc")
                if loc and loc.get_text(strip=True) not in visited:
                    pending.append(loc.get_text(strip=True))
            for node in soup.find_all("url"):
                loc = node.find("loc")
                if loc:
                    u = normalise_url(loc.get_text(strip=True))
                    if u not in urls:
                        urls.append(u)
                        if len(urls) >= max_urls:
                            break
        except Exception:
            pass
    return urls


def crawl_site(root_url: str, max_pages: int, delay: float, ignore_robots: bool, include_subdomains: bool, workers: int, progress_callback=None) -> dict[str, Page]:
    """Breadth-first same-domain crawl, scheduling requests in bounded batches."""
    root_url = normalise_url(root_url)
    robots = load_robots(root_url, ignore_robots)
    seen, pages = set(), {}
    queued = deque([root_url])
    queued_set = {root_url}

    # Pre-seed queue from sitemap for high page discovery
    try:
        sitemap_urls = extract_sitemap_urls(root_url, max_urls=max_pages)
        for u in sitemap_urls:
            if u not in queued_set and same_domain(u, root_url, include_subdomains):
                queued.append(u)
                queued_set.add(u)
    except Exception:
        pass

    with ThreadPoolExecutor(max_workers=workers) as pool, tqdm(total=max_pages, desc="Crawling", unit="page", disable=not safe_is_tty()) as progress:
        while queued and len(seen) < max_pages:
            batch = []
            while queued and len(seen) + len(batch) < max_pages and len(batch) < workers:
                url = queued.popleft()
                if url in seen or not allowed(url, robots):
                    continue
                seen.add(url)
                batch.append(url)
            if not batch:
                continue
            futures = {pool.submit(analyse_page, u): u for u in batch}
            for future in as_completed(futures):
                page = future.result()
                pages[page.url] = page
                progress.update(1)
                if progress_callback:
                    try:
                        progress_callback(len(pages), max_pages, page.url)
                    except Exception:
                        pass
                for link in page.internal_links:
                    if same_domain(link.url, root_url, include_subdomains) and link.url not in seen and link.url not in queued_set:
                        queued.append(link.url)
                        queued_set.add(link.url)
            if delay:
                time.sleep(delay)  # Delay between request batches, not each concurrent request.
    return pages


def check_url(url: str) -> tuple[int | None, int | None, str]:
    try:
        response = fetch(url, stream=True)
        size = int(response.headers.get("Content-Length", 0))
        response.close()
        return response.status_code, size, response.url
    except requests.RequestException as exc:
        logging.error("Resource %s: %s", url, exc)
        return None, None, ""


def find_broken_links(pages: dict[str, Page], root: str, include_subdomains: bool, workers: int):
    links, images, assets = {}, {}, {}
    for page in pages.values():
        for item in page.internal_links:
            if same_domain(item.url, root, include_subdomains): links[(page.url, item.url, item.anchor)] = item
        for item in page.images: images[(page.url, item.url, item.alt)] = item
        for item in page.assets: assets[(page.url, item.url, item.alt)] = item
    statuses = {}
    unique = {key[1] for key in links} | {key[1] for key in images} | {key[1] for key in assets}
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {pool.submit(check_url, url): url for url in unique}
        for future in tqdm(as_completed(futures), total=len(futures), desc="Checking resources", unit="file", disable=not safe_is_tty()):
            statuses[futures[future]] = future.result()
    broken_links = [[s, u, statuses[u][0] or "Request failed", a] for s, u, a in links if not statuses[u][0] or statuses[u][0] >= 400]
    broken_images = [[s, u, statuses[u][0] or "Request failed", a] for s, u, a in images if not statuses[u][0] or statuses[u][0] >= 400]
    unminified = []
    for source, url, kind in assets:
        status, size, _ = statuses[url]
        filename = urlparse(url).path.lower()
        if status and status < 400 and ".min." not in filename and not filename.endswith(".min"):
            unminified.append([url, kind, round((size or 0) / 1024, 1), source, "No .min. in filename (review minification)"])
    return broken_links, broken_images, unminified


def find_duplicates(pages: dict[str, Page]):
    title_groups, content_groups = defaultdict(list), defaultdict(list)
    for p in pages.values():
        if p.title: title_groups[p.title].append(p.url)
        if p.content_hash and p.status and p.status < 400: content_groups[p.content_hash].append(p)
    titles = [[title, ", ".join(urls), len(urls)] for title, urls in title_groups.items() if len(urls) > 1]
    content = []
    for digest, group in content_groups.items():
        if len(group) > 1:
            content.append([digest, group[0].url, group[1].url, "100%", group[0].word_count])
    return titles, content


def parse_sitemap(root: str, pages: dict[str, Page], ignore_robots: bool, workers: int) -> list[list]:
    """Read sitemap.xml and sitemap indexes, then report audit issues for listed URLs."""
    pending, visited, entries = deque([urljoin(root, "/sitemap.xml")]), set(), []
    while pending:
        sitemap = pending.popleft()
        if sitemap in visited: continue
        visited.add(sitemap)
        try:
            r = fetch(sitemap)
            if r.status_code >= 400: entries.append([sitemap, "", "404", r.status_code]); continue
            soup = BeautifulSoup(r.content, "xml")
            for node in soup.find_all("sitemap"):
                loc = node.find("loc")
                if loc: pending.append(loc.get_text(strip=True))
            for node in soup.find_all("url"):
                loc = node.find("loc")
                if loc: entries.append([sitemap, normalise_url(loc.get_text(strip=True)), "", ""])
        except requests.RequestException as exc:
            logging.error("Sitemap %s: %s", sitemap, exc)
            entries.append([sitemap, "", "404", "Request failed"])
    for row in entries:
        if not row[1] or row[2]: continue
        # The crawler has already fetched in-scope pages. Reusing that result
        # avoids a second, potentially massive crawl of every sitemap entry.
        page = pages.get(row[1])
        if not page:
            row[2], row[3] = "Missing from Crawl", "Not crawled"
        elif not page.status or page.status >= 400:
            row[2], row[3] = "404", page.status or "Request failed"
        elif page.redirect_chain:
            row[2], row[3] = "Redirected", page.status
        elif page.indexability == "Noindex":
            row[2], row[3] = "Noindex", page.status
        elif page.canonical and page.canonical != page.url:
            row[2], row[3] = "Non-Canonical", page.status
        else: continue
    return [r for r in entries if r[2]]


def generate_excel_report(output: Path, pages: dict[str, Page], broken_links, broken_images, duplicate_titles, duplicate_content, unminified, sitemap_issues, threshold: int, started: datetime):
    wb = Workbook(); wb.remove(wb.active)
    pages_list = sorted(pages.values(), key=lambda p: p.url)
    on_page, thin, h1s, underscores = [], [], [], []
    for p in pages_list:
        issues = []
        if not p.title: issues.append("Missing title")
        if not p.meta_description: issues.append("Missing meta description")
        if not p.h1s: issues.append("Missing H1")
        if p.word_count < threshold: issues.append(f"Thin content (<{threshold})")
        if issues: on_page.append([p.url, "Y" if not p.title else "", "Y" if not p.meta_description else "", "Y" if not p.h1s else "", "Y" if p.word_count < threshold else "", "; ".join(issues)])
        if p.word_count < threshold: thin.append([p.url, p.word_count, p.title])
        if len(p.h1s) > 1: h1s.append([p.url, len(p.h1s), " | ".join(p.h1s)])
        if "_" in urlparse(p.url).path: underscores.append([p.url, p.url.replace("_", "-")])
    errors = sum(1 for p in pages_list if p.error or (p.status or 0) >= 400)
    duration = round((datetime.now() - started).total_seconds(), 1)
    summary = [["Metric", "Value"], ["Total URLs crawled", len(pages_list)], ["Total errors", errors], ["Broken links count", len(broken_links)], ["Broken images count", len(broken_images)], ["Duplicate titles count", len(duplicate_titles)], ["Duplicate content count", len(duplicate_content)], ["Thin content pages count", len(thin)], ["Multiple H1 pages count", len(h1s)], ["Underscore URLs count", len(underscores)], ["Unminified CSS/JS count", len(unminified)], ["Sitemap issues count", len(sitemap_issues)], ["Avg word count", round(sum(p.word_count for p in pages_list) / len(pages_list), 1) if pages_list else 0], ["Crawl date/time", started.strftime("%Y-%m-%d %H:%M:%S")], ["Crawl duration (seconds)", duration]]
    sheets = {
        "Summary": summary, "All Urls": [["URL", "Title", "Title Length", "Meta Description", "Word Count", "HTTP Status", "Response Time (ms)", "Indexability"]] + [[p.url, p.title, len(p.title), p.meta_description, p.word_count, p.status or "Request failed", p.response_ms or "", p.indexability] for p in pages_list],
        "On Page Errors": [["URL", "Missing Title", "Missing Meta Description", "Missing H1", "Thin Content (Y/N)", "Issue Details"]] + on_page,
        "Broken Internal Links": [["Source Page", "Broken Link URL", "HTTP Status", "Anchor Text"]] + broken_links,
        "Broken Images": [["Source Page", "Image URL", "HTTP Status", "Alt Text"]] + broken_images,
        "Duplicate Titles": [["Title", "URLs (comma separated)", "Count"]] + duplicate_titles,
        "Duplicate Content": [["Content Hash", "URL 1", "URL 2", "Similarity %", "Word Count"]] + duplicate_content,
        "Thin Content Pages": [["URL", "Word Count", "Title"]] + thin,
        "Multiple H1 Tags": [["URL", "H1 Count", "H1 Texts (pipe-separated)"]] + h1s,
        "Underscore Urls": [["URL", "Suggested URL"]] + underscores,
        "Unminified CSS & JS": [["File URL", "Type (CSS/JS)", "File Size (KB)", "Found On Page", "Reason"]] + unminified,
        "Sitemap Issues": [["Sitemap URL", "Page URL", "Issue Type", "HTTP Status"]] + sitemap_issues,
    }
    header_fill, red_fill = PatternFill("solid", fgColor="1F4E78"), PatternFill("solid", fgColor="FCE4D6")
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows: ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = "A1:" + get_column_letter(max(1, ws.max_column)) + str(max(1, ws.max_row))
        for cell in ws[1]: cell.font, cell.fill = Font(bold=True, color="FFFFFF"), header_fill
        for col in range(1, ws.max_column + 1):
            width = min(60, max(12, max((len(str(ws.cell(r, col).value or "")) for r in range(1, ws.max_row + 1)), default=10) + 2))
            ws.column_dimensions[get_column_letter(col)].width = width
        # Findings-only sheets use a light-red visual cue for populated issue cells.
        if name in {"On Page Errors", "Broken Internal Links", "Broken Images", "Sitemap Issues"} and ws.max_row > 1:
            ws.conditional_formatting.add(f"A2:{get_column_letter(ws.max_column)}{ws.max_row}", CellIsRule(operator="notEqual", formula=['""'], fill=red_fill))
        if name == "All Urls" and ws.max_row > 1:
            ws.conditional_formatting.add(f"F2:F{ws.max_row}", CellIsRule(operator="equal", formula=["200"], fill=PatternFill("solid", fgColor="C6EFCE")))
            ws.conditional_formatting.add(f"F2:F{ws.max_row}", CellIsRule(operator="greaterThanOrEqual", formula=["400"], fill=red_fill))
    wb.save(output)
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description="Crawl a website and produce a 12-sheet SEO Excel audit.")
    parser.add_argument("--url", required=True, help="Root URL, e.g. https://example.com")
    parser.add_argument("--max-pages", type=int, default=500)
    parser.add_argument("--output", help="Output .xlsx path (default: SEO_Audit_<domain>_<date>.xlsx)")
    parser.add_argument("--delay", type=float, default=0.5, help="Seconds between batches")
    parser.add_argument("--ignore-robots", action="store_true")
    parser.add_argument("--thin-threshold", type=int, default=150)
    parser.add_argument("--include-subdomains", action="store_true")
    parser.add_argument("--workers", type=int, default=8, choices=range(1, 33))
    args = parser.parse_args()
    if not urlparse(args.url).scheme: args.url = "https://" + args.url
    logging.basicConfig(filename="crawl_errors.log", level=logging.ERROR, format="%(asctime)s %(levelname)s %(message)s")
    started = datetime.now()
    root = normalise_url(args.url)
    pages = crawl_site(root, args.max_pages, args.delay, args.ignore_robots, args.include_subdomains, args.workers)
    broken_links, broken_images, unminified = find_broken_links(pages, root, args.include_subdomains, args.workers)
    duplicate_titles, duplicate_content = find_duplicates(pages)
    sitemap_issues = parse_sitemap(root, pages, args.ignore_robots, args.workers)
    domain = re.sub(r"[^A-Za-z0-9.-]+", "_", urlparse(root).netloc)
    output = Path(args.output or f"SEO_Audit_{domain}_{started:%Y%m%d}.xlsx")
    summary = generate_excel_report(output, pages, broken_links, broken_images, duplicate_titles, duplicate_content, unminified, sitemap_issues, args.thin_threshold, started)
    print("\nSEO audit complete:", output.resolve())
    for metric, value in summary[1:]: print(f"  {metric}: {value}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
